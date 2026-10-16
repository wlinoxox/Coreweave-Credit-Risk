"""Post-process the built deck: larger table text, content-based column widths, no auto chart titles."""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.chart import XL_TICK_LABEL_POSITION


def text_width(s):
    # CJK characters are roughly twice as wide as Latin ones
    return sum(2 if ord(c) > 0x2E80 else 1 for c in s)


def polish(path):
    prs = Presentation(path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                chart = shape.chart
                chart.has_title = False
                # keep category labels below the plot when bars go negative
                if any(v is not None and v < 0 for s in chart.plots[0].series for v in s.values):
                    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            if not shape.has_table:
                continue
            tbl = shape.table
            ncols = len(tbl.columns)
            weights = [max(text_width(tbl.cell(r, c).text) for r in range(len(tbl.rows))) + 4
                       for c in range(ncols)]
            total = shape.width
            for c, w in enumerate(weights):
                tbl.columns[c].width = Emu(int(total * w / sum(weights)))
            dense = len(tbl.rows) > 8
            body_pt = 12 if dense else 13 if ncols > 5 else 14 if ncols > 2 else 15
            row_pt = 28 if dense else 34
            for r, row in enumerate(tbl.rows):
                row.height = Pt(row_pt)
                for cell in row.cells:
                    cell.vertical_anchor = 3  # middle
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(body_pt)
    prs.save(path)


if __name__ == "__main__":
    polish(sys.argv[1])
    print("polished", sys.argv[1])
